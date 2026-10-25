#!/usr/bin/env python3
"""Recalcule un .xlsx avec LibreOffice (headless) puis signale les erreurs de formule.
Usage: python3 recalc.py fichier.xlsx [timeout_s]   (nécessite 'soffice' dans le PATH)"""
import json, shutil, subprocess, sys, tempfile, os
from openpyxl import load_workbook
path = os.path.abspath(sys.argv[1]); timeout = int(sys.argv[2]) if len(sys.argv) > 2 else 90
soffice = shutil.which("soffice") or shutil.which("libreoffice")
if not soffice: print(json.dumps({"error": "LibreOffice (soffice) introuvable"})); sys.exit(1)
with tempfile.TemporaryDirectory() as d:
    subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", d, path],
                   check=True, timeout=timeout, capture_output=True)
    shutil.copy(os.path.join(d, os.path.basename(path)), path)
wb = load_workbook(path, data_only=True); errs = {}; total = 0
wf = load_workbook(path)
for ws in wf:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                total += 1; v = wb[ws.title][c.coordinate].value
                if isinstance(v, str) and v.startswith("#"): errs.setdefault(v, []).append(f"{ws.title}!{c.coordinate}")
print(json.dumps({"status": "errors_found" if errs else "success", "total_formulas": total,
                  "total_errors": sum(map(len, errs.values())), "error_summary": errs}, ensure_ascii=False, indent=2))
