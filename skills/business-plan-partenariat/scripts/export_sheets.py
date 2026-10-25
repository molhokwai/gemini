# -*- coding: utf-8 -*-
"""Exporte des feuilles du classeur recalculé en blocs (titre, intertitre, note, tableau) pour le .docx.

Usage: python3 export_sheets.py classeur.xlsx sortie.plan.json "Résumé,Achats-Charges,..."
Lit les valeurs calculées (le classeur doit avoir été recalculé: recalc.py) et les styles posés par
build_xlsx.py: bandeau de couleur fusionné = titre/intertitre, police grise italique = note,
fond gris BBBBBB = en-tête de tableau. Les montants sont rendus avec le format de la cellule.
"""
import json, sys
from openpyxl import load_workbook

path, out, names = sys.argv[1], sys.argv[2], [n.strip() for n in sys.argv[3].split(",") if n.strip()]
wb = load_workbook(path, data_only=True)

def rgb(c):
    try:
        v = c.fill.fgColor.rgb if c.fill and c.fill.fill_type == "solid" else None
        return v[-6:].upper() if isinstance(v, str) else None
    except Exception: return None
def font_rgb(c):
    try:
        v = c.font.color.rgb if c.font and c.font.color else None
        return v[-6:].upper() if isinstance(v, str) else None
    except Exception: return None

def fmt(c):
    v, f = c.value, c.number_format or ""
    if v is None: return ""
    if isinstance(v, bool): return "Oui" if v else "Non"
    if isinstance(v, (int, float)):
        if "%" in f:
            s = f"{v * 100:.1f}".replace(".", ",")
            return (s[:-2] if s.endswith(",0") else s) + " %"
        if '"' in f:  # devise: '#,##0" f"'
            suffix = f.split('"')[1]
            return f"{round(v):,}".replace(",", " ") + suffix.replace(" ", " ")
        if f in ("0.0",): return f"{v:.1f}".replace(".", ",")
        if float(v).is_integer(): return f"{int(v):,}".replace(",", " ")
        return f"{v:.2f}".rstrip("0").rstrip(",.").replace(".", ",")
    return str(v)

sheets = []
for name in names:
    if name not in wb.sheetnames:
        print(f"ATTENTION: feuille absente: {name}", file=sys.stderr); continue
    ws = wb[name]
    merged = {}  # coin haut-gauche -> (dernière colonne, dernière ligne)
    covered = set()
    for rng in ws.merged_cells.ranges:
        merged[(rng.min_row, rng.min_col)] = (rng.max_col, rng.max_row)
        for r in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                if (r, col) != (rng.min_row, rng.min_col): covered.add((r, col))
    hidden_rows = {r for r, d in ws.row_dimensions.items() if d.hidden}
    hidden_cols = {k for k, d in ws.column_dimensions.items() if d.hidden}
    maxc = ws.max_column
    from openpyxl.utils import get_column_letter
    cols = [c for c in range(1, maxc + 1) if get_column_letter(c) not in hidden_cols]
    blocks, table = [], None
    def flush():
        global table
        if table and table["rows"]:
            # retire les colonnes entièrement vides
            keep = [i for i in range(len(table["rows"][0]["cells"])) if any(r["cells"][i]["t"] for r in table["rows"])]
            for r in table["rows"]: r["cells"] = [r["cells"][i] for i in keep]
            table["widths"] = [table["widths"][i] for i in keep]
            blocks.append(table)
        table = None
    for r in range(1, ws.max_row + 1):
        if r in hidden_rows or (r, 1) in covered: continue
        a = ws.cell(r, 1)
        if (r, 1) in merged and merged[(r, 1)][0] >= 3:  # bandeau ou note sur toute la largeur
            if a.value is None: continue
            flush()
            if font_rgb(a) == "808080" or (a.font and a.font.italic):
                blocks.append({"type": "note", "text": str(a.value)})
            else:
                blocks.append({"type": "title" if r == 1 else "heading", "text": str(a.value)})
            continue
        vals = [ws.cell(r, c) for c in cols]
        if all(v.value in (None, "") for v in vals):
            flush(); continue
        filled = [v for v in vals if v.value not in (None, "")]
        if len(filled) == 1 and filled[0].column == 1 and (font_rgb(a) == "808080" or (a.font and a.font.italic)):
            flush(); blocks.append({"type": "note", "text": str(a.value)}); continue
        header = rgb(a) == "BBBBBB" or all(rgb(v) == "BBBBBB" for v in filled)
        if table is None:
            table = {"type": "table", "rows": [],
                     "widths": [ws.column_dimensions[get_column_letter(c)].width or 10 for c in cols]}
        cells = []
        for v in vals:
            cells.append({"t": "" if (v.row, v.column) in covered else fmt(v),
                          "fill": rgb(v), "color": font_rgb(v), "bold": bool(v.font and v.font.bold),
                          "italic": bool(v.font and v.font.italic), "num": isinstance(v.value, (int, float))})
        table["rows"].append({"header": header, "cells": cells})
    flush()
    ncols = max([len(b["rows"][0]["cells"]) for b in blocks if b["type"] == "table"] or [0])
    sheets.append({"name": name, "blocks": blocks, "max_cols": ncols})

json.dump({"sheets": sheets}, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("OK", out, [(s["name"], len(s["blocks"]), s["max_cols"]) for s in sheets])
